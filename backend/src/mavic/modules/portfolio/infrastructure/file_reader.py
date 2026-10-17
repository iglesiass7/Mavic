"""Lectura de archivos CSV y XLSX a una lista de filas.

Solo LEE: convierte bytes en filas `{columna: valor}`. No valida el negocio
(eso lo hace `domain.import_rows`).

Defensas contra archivos maliciosos o enormes:
- Tamaño máximo del archivo subido (lo controla el caso de uso).
- Máximo de filas: se corta la lectura apenas se pasa el límite.
- XLSX "bomba zip": un .xlsx es un ZIP. Un archivo de 1 MB puede descomprimirse
  a varios GB y tirar el servidor. Antes de abrirlo se suma el tamaño
  descomprimido declarado de sus partes y se rechaza si es demasiado grande.
"""

import csv
import io
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass

from openpyxl import load_workbook

from mavic.core.errors import BusinessRuleViolationError

MAX_ROWS = 5000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
SUPPORTED_EXTENSIONS = (".csv", ".xlsx")

RawRow = tuple[int, dict[str, object]]  # (número de fila en la planilla, {columna: valor})


@dataclass(frozen=True)
class Table:
    headers: list[str]
    rows: list[RawRow]


def _error(code: str, message: str) -> BusinessRuleViolationError:
    return BusinessRuleViolationError(message, code=code)


def _normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip().lower()


def _is_empty(values: Iterable[object]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def _build_table(lines: Iterator[Iterable[object]]) -> Table:
    """Convierte una secuencia de filas (la primera es el encabezado) en una Table."""
    header_line = next(lines, None)
    if header_line is None:
        raise _error("EMPTY_FILE", "El archivo está vacío.")
    headers = [_normalize_header(value) for value in header_line]

    rows: list[RawRow] = []
    for row_number, line in enumerate(lines, start=2):  # la fila 1 es el encabezado
        cells = list(line)
        if _is_empty(cells):
            continue
        if len(rows) >= MAX_ROWS:
            raise _error("TOO_MANY_ROWS", f"El archivo supera el máximo de {MAX_ROWS} filas.")
        rows.append(
            (row_number, {h: cells[i] if i < len(cells) else None for i, h in enumerate(headers)})
        )
    if not rows:
        raise _error("EMPTY_FILE", "El archivo no tiene filas con datos.")
    return Table(headers=headers, rows=rows)


def _decode(content: bytes) -> str:
    # utf-8-sig ignora la "marca BOM" que agrega Excel al exportar en UTF-8.
    # Si no es UTF-8, se prueba cp1252: la codificación de Excel en Windows en español.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise _error("UNREADABLE_FILE", "No se pudo leer el texto del archivo CSV.")


def _read_csv(content: bytes) -> Table:
    text = _decode(content)
    try:
        # Excel en español exporta CSV separado por ";", no por ",". Se detecta solo.
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    return _build_table(iter(csv.reader(io.StringIO(text), delimiter=delimiter)))


def _check_zip_bomb(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            uncompressed = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise _error("UNREADABLE_FILE", "El archivo .xlsx está dañado o no es un Excel.") from exc
    if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise _error("FILE_TOO_LARGE", "El contenido del Excel es demasiado grande.")


def _read_xlsx(content: bytes) -> Table:
    _check_zip_bomb(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza tipos muy variados ante archivos rotos
        raise _error("UNREADABLE_FILE", "No se pudo abrir el archivo Excel.") from exc
    try:
        sheet = workbook.worksheets[0]  # siempre la PRIMERA hoja, no "la última abierta"
        return _build_table(iter(sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()


def read_table(filename: str, content: bytes) -> Table:
    name = filename.lower()
    if name.endswith(".csv"):
        return _read_csv(content)
    if name.endswith(".xlsx"):
        return _read_xlsx(content)
    raise _error(
        "UNSUPPORTED_FILE_TYPE",
        "Formato no soportado. Subí un archivo .csv o .xlsx (Excel 2007 o posterior).",
    )
